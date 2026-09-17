"""Text- und Strukturextraktion aus Originaldokumenten.

``extract(document, data)`` wählt anhand der Dateiendung das passende Verfahren und liefert
Segmente: eine PDF-Seite, ein Abschnitt unter einer Überschrift, eine Tabellenzeile.

Reine Funktionen: Sie bekommen den Dateiinhalt als Bytes und greifen nicht auf das
Dateisystem zu. Das macht sie ohne Testdateien auf der Platte prüfbar und hält den
Zugriff auf die Originale in ``pipeline`` an einer Stelle.

Leitgedanke bei Tabellen: Eine Zeile bleibt **ein** Segment und wird als
``Spaltenname: Wert`` geschrieben. Zerfiele sie in einzelne Zellen, verlöre man den Bezug
zwischen Thema, Sachstand und Verantwortlichem – und damit die Antwort auf Fragen wie
„Wer ist verantwortlich?“.
"""

import io
from collections.abc import Iterator
from datetime import date, datetime
from pathlib import Path

import docx
import openpyxl
import pymupdf
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from bauprojekt.models import Document, Segment, SegmentKind, normalize_compatibility

MAX_HEADING_LENGTH = 80
"""Längerer Text ist Fließtext, keine Überschrift."""

MIN_HEADING_SIZE_RATIO = 1.15
"""So viel größer als der Fließtext muss eine Zeile sein, um als Überschrift zu gelten."""

MIN_HEADER_CELLS = 3
"""Ab so vielen gefüllten Zellen gilt eine Zeile im Tabellenblatt als Kopfzeile."""


def extract(document: Document, data: bytes) -> list[Segment]:
    """Dateiinhalt in Segmente zerlegen. Neue Dateiarten kommen hier als weiterer Fall dazu."""
    suffix = Path(document.file_name).suffix.lower()
    match suffix:
        case ".pdf":
            return extract_pdf(document, data)
        case ".docx":
            return extract_docx(document, data)
        case ".xlsx":
            return extract_xlsx(document, data)
        case _:
            raise ValueError(f"Nicht unterstützte Dateiendung: {suffix}")


def extract_pdf(document: Document, data: bytes) -> list[Segment]:
    """Eine Seite wird ein Segment – so bleibt die Seitenzahl als Quellenangabe erhalten."""
    segments: list[Segment] = []
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        for page_no, page in enumerate(pdf, start=1):
            text = normalize_compatibility(page.get_text("text")).strip()
            if not text:
                continue  # Leere oder rein grafische Seiten übergehen.
            segments.append(
                Segment.create(
                    document=document,
                    index=len(segments),
                    kind=SegmentKind.SEITE,
                    page_no=page_no,
                    locator=f"S. {page_no}",
                    heading=heading_by_font_size(page),
                    text=text,
                )
            )
    return segments


def extract_docx(document: Document, data: bytes) -> list[Segment]:
    """Absätze werden unter ihrer Überschrift gesammelt, Tabellenzeilen einzeln abgelegt."""
    docx_document = docx.Document(io.BytesIO(data))
    segments: list[Segment] = []
    heading: str | None = None
    paragraphs: list[str] = []
    table_no = 0
    section_no = 0

    def flush_paragraphs() -> None:
        """Gesammelte Absätze als ein Abschnitt ablegen."""
        nonlocal paragraphs, section_no
        text = "\n".join(paragraphs).strip()
        if text:
            section_no += 1
            segments.append(
                Segment.create(
                    document=document,
                    index=len(segments),
                    kind=SegmentKind.ABSCHNITT,
                    locator=heading or f"Abschnitt {section_no}",
                    heading=heading,
                    text=text,
                )
            )
        paragraphs = []

    for block in iter_blocks(docx_document):
        if isinstance(block, Paragraph):
            if is_heading(block):
                flush_paragraphs()
                heading = block.text.strip()
            elif block.text.strip():
                paragraphs.append(block.text.strip())
            continue

        flush_paragraphs()
        table_no += 1
        # Nach der Tabelle gilt die Überschrift davor nicht weiter: Ein Schlussabsatz
        # gehört nicht mehr unter „Tagesordnungspunkte“.
        heading_before_table, heading = heading, None
        segments.extend(
            table_row_segments(
                document=document,
                rows=[[cell.text.strip() for cell in row.cells] for row in block.rows],
                heading=heading_before_table,
                locator_prefix=f"Tabelle {table_no}",
                first_index=len(segments),
            )
        )

    flush_paragraphs()
    return segments


def extract_xlsx(document: Document, data: bytes) -> list[Segment]:
    """Je Tabellenblatt: Zeilen über der Kopfzeile als Abschnitt, danach eine Zeile je Segment."""
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    segments: list[Segment] = []
    for sheet in workbook.worksheets:
        rows = [
            [format_cell(value) for value in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        header_index = find_header_row(rows)
        if header_index is None:
            continue

        preamble = "\n".join(
            " ".join(cell for cell in row if cell).strip()
            for row in rows[:header_index]
        )
        if preamble.strip():
            segments.append(
                Segment.create(
                    document=document,
                    index=len(segments),
                    kind=SegmentKind.ABSCHNITT,
                    locator=f"{sheet.title}, Kopf",
                    heading=sheet.title,
                    text=preamble.strip(),
                )
            )

        segments.extend(
            table_row_segments(
                document=document,
                rows=rows[header_index:],
                heading=sheet.title,
                locator_prefix=sheet.title,
                first_index=len(segments),
            )
        )
    return segments


# --------------------------------------------------------------------------- Hilfsfunktionen


def iter_blocks(document: DocxDocument) -> Iterator[Paragraph | Table]:
    """Absätze und Tabellen in Dokumentreihenfolge durchlaufen.

    python-docx bietet ``paragraphs`` und ``tables`` getrennt an; die Reihenfolge im
    Dokument ergibt sich erst aus dem XML.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, document)
        elif child.tag == qn("w:tbl"):
            yield Table(child, document)


def is_heading(paragraph: Paragraph) -> bool:
    style = (paragraph.style.name or "") if paragraph.style else ""
    return style.startswith(("Heading", "Überschrift", "Title")) and bool(
        paragraph.text.strip()
    )


def heading_by_font_size(page: pymupdf.Page) -> str | None:
    """Die Zeile mit der größten Schrift als Überschrift deuten.

    Die erste Zeile taugt nicht: Bei Behördenschreiben steht dort der Briefkopf. Die
    Schriftgröße trennt Überschrift und Fließtext zuverlässiger, weil sie die Auszeichnung
    im Dokument selbst nutzt.
    """
    lines: list[tuple[float, str]] = []
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            size = max((span["size"] for span in line["spans"]), default=0.0)
            text = normalize_compatibility(
                "".join(span["text"] for span in line["spans"])
            ).strip()
            if text:
                lines.append((size, text))

    if not lines:
        return None

    largest_size, largest_text = max(lines, key=lambda entry: entry[0])
    smaller_sizes = [size for size, _ in lines if size < largest_size]
    if not smaller_sizes:
        return None  # Alles gleich groß: keine Auszeichnung, also keine Überschrift.

    body_size = max(smaller_sizes)
    if (
        largest_size < body_size * MIN_HEADING_SIZE_RATIO
        or len(largest_text) > MAX_HEADING_LENGTH
    ):
        return None
    return largest_text


def format_cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().strftime("%d.%m.%Y")
    if isinstance(value, date):
        return value.strftime("%d.%m.%Y")
    return str(value).strip()


def find_header_row(rows: list[list[str]]) -> int | None:
    """Erste Zeile mit genügend gefüllten Zellen. Darüber stehen Titel und Leerzeilen."""
    for index, row in enumerate(rows):
        if sum(1 for cell in row if cell) >= MIN_HEADER_CELLS:
            return index
    return None


def table_row_segments(
    *,
    document: Document,
    rows: list[list[str]],
    heading: str | None,
    locator_prefix: str,
    first_index: int,
) -> list[Segment]:
    """Tabellenzeilen ab der Kopfzeile in Segmente überführen (``Spaltenname: Wert``)."""
    if not rows:
        return []

    header = rows[0]
    numbered_first_column = header[0].strip().upper() == "TOP" if header else False
    segments: list[Segment] = []

    for row_no, row in enumerate(rows[1:], start=1):
        values = [
            f"{label}: {value}"
            for label, value in zip(header, row, strict=False)
            if value and label
        ]
        if not values:
            continue

        key = row[0].strip() if row and row[0].strip() else str(row_no)
        locator = (
            f"TOP {key}" if numbered_first_column else f"{locator_prefix}, Zeile {key}"
        )
        segments.append(
            Segment.create(
                document=document,
                index=first_index + len(segments),
                kind=SegmentKind.TABELLENZEILE,
                locator=locator,
                heading=heading,
                text="\n".join(values),
            )
        )
    return segments
