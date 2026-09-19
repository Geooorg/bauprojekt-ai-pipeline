"""Gemeinsame Fixtures: kleine Beispieldateien, im Test erzeugt statt eingecheckt.

Sie bilden die Struktur der echten Testdaten nach (PDF mit Seiten, Protokoll mit TOP-Tabelle,
Terminplan mit Kopfzeile), bleiben aber winzig, damit die Tests schnell und lesbar sind.
"""

import io
from collections.abc import Iterator
from datetime import UTC, date, datetime

import docx
import openpyxl
import psycopg
import pymupdf
import pytest

from bauprojekt.config import DATABASE_URL, EMBEDDING_DIM
from bauprojekt.db import init_schema
from bauprojekt.models import DocType, Document, derive_id

INGESTED_AT = datetime(2026, 9, 17, 8, 0, tzinfo=UTC).replace(tzinfo=None)


def make_document(
    *,
    file_name: str,
    doc_type: DocType,
    media_type: str,
    project_id: str = "BAU-42",
    document_date: date | None = date(2026, 9, 15),
) -> Document:
    """Document ohne Datei auf der Platte – die Extraktion arbeitet auf Bytes."""
    source_path = f"{project_id}/ordner/{file_name}"
    content_hash = "b" * 64
    return Document(
        document_id=derive_id(project_id, source_path, content_hash),
        content_hash=content_hash,
        project_id=project_id,
        source_path=source_path,
        file_name=file_name,
        doc_type=doc_type,
        media_type=media_type,
        file_size=1234,
        document_date=document_date,
        ingested_at=INGESTED_AT,
    )


@pytest.fixture
def pdf_bytes() -> bytes:
    """Zweiseitiges PDF mit Überschrift und Fließtext."""
    doc = pymupdf.open()
    seiten = [
        "<h1>Projektstatusbericht September 2026</h1><p>Gesamtstatus: ROT. Auflagen offen.</p>",
        (
            "<p style='font-size:8pt'>Stadt Musterstadt · Bauaufsichtsamt · Rathausplatz 1</p>"
            "<h1>2. Termine</h1><p>Baugenehmigung Haus B: Antrag ruht.</p>"
        ),
    ]
    for html in seiten:
        page = doc.new_page(width=595, height=842)
        page.insert_htmlbox(
            page.rect + (56, 56, -56, -56), html, css="h1 { font-size: 16pt; }"
        )
    return doc.tobytes()


@pytest.fixture
def docx_bytes() -> bytes:
    """Protokoll mit Überschriften, Absätzen und TOP-Tabelle."""
    document = docx.Document()
    document.add_heading("Protokoll Baubesprechung Nr. 25", level=1)
    document.add_paragraph("Projekt: BAU-42")
    document.add_heading("Tagesordnungspunkte", level=2)
    table = document.add_table(rows=1, cols=4)
    for cell, text in zip(
        table.rows[0].cells,
        ("TOP", "Thema", "Sachstand", "Verantwortlich"),
        strict=True,
    ):
        cell.text = text
    zeilen = [
        ("1", "Rohbau Haus A", "Verzug unverändert 2 Wochen.", "Kessler Bau"),
        ("2", "Brandschutzprüfung", "Kein Verantwortlicher benannt.", "ungeklärt"),
    ]
    for zeile in zeilen:
        for cell, text in zip(table.add_row().cells, zeile, strict=True):
            cell.text = text
    document.add_paragraph("Nächste Baubesprechung: 22.09.2026, 10:00 Uhr")
    puffer = io.BytesIO()
    document.save(puffer)
    return puffer.getvalue()


@pytest.fixture
def xlsx_bytes() -> bytes:
    """Terminplan mit Titelzeile, Leerzeile und Kopfzeile."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Bauzeitenplan"
    sheet.append(["Bauzeitenplan BAU-42 – Stand 15.09.2026", None, None, None])
    sheet.append([])
    sheet.append(["ID", "Vorgang", "Verzug (AT)", "Status"])
    sheet.append(["V02", "Rohbau Haus A", 10, "verzögert"])
    sheet.append(["V07", "Baugenehmigung Haus B", None, "offen"])
    puffer = io.BytesIO()
    workbook.save(puffer)
    return puffer.getvalue()


@pytest.fixture
def db() -> Iterator[psycopg.Connection]:
    """Verbindung auf ein eigenes Schema, das nach jedem Test wieder verschwindet.

    So laufen die Tests gegen dieselbe Datenbank wie die Anwendung, ohne deren Daten
    anzufassen. Ist keine Datenbank erreichbar, wird der Test übersprungen.
    """
    try:
        connection = psycopg.connect(DATABASE_URL, autocommit=True)
    except psycopg.OperationalError as fehler:
        pytest.skip(f"Keine Datenbank erreichbar: {fehler}")

    with connection:
        connection.execute("CREATE SCHEMA IF NOT EXISTS pytest_bauprojekt")
        connection.execute("SET search_path TO pytest_bauprojekt, public")
        init_schema(connection)
        try:
            yield connection
        finally:
            connection.execute("DROP SCHEMA pytest_bauprojekt CASCADE")


class FakeEncoder:
    """Doppel für das Embedding-Modell: merkt sich die Eingaben, braucht weder Modell noch Torch.

    Ohne ``vectors`` liefert es für jeden Text einen Vektor aus seiner Länge. Mit ``vectors``
    lässt sich gezielt steuern, welcher Text wohin im Vektorraum fällt – nützlich, um die
    Reihenfolge von Suchtreffern vorherzusagen.
    """

    def __init__(
        self,
        dimension: int = EMBEDDING_DIM,
        vectors: dict[str, list[float]] | None = None,
    ) -> None:
        self.dimension = dimension
        self.vectors = vectors or {}
        self.calls: list[list[str]] = []

    def encode(self, texts: list[str]) -> list[list[float]]:
        self.calls.append(list(texts))
        return [self.vector_for(text) for text in texts]

    def vector_for(self, text: str) -> list[float]:
        for fragment, vector in self.vectors.items():
            if fragment in text:
                return vector
        return [float(len(text))] * self.dimension

    @property
    def eingaben(self) -> list[str]:
        return [text for call in self.calls for text in call]


def basis(position: int, dimension: int = EMBEDDING_DIM) -> list[float]:
    """Einheitsvektor: 1.0 an einer Stelle, sonst 0. Zwei verschiedene sind orthogonal."""
    vector = [0.0] * dimension
    vector[position] = 1.0
    return vector
