"""Erzeugt synthetische Projektdokumente für BAU-42 und BAU-43 unter data/raw/.

Die Dokumente enthalten bewusst bekannte Sachverhalte (siehe docs/testdaten.md),
damit Extraktion, Suche und Risikoanalyse gegen eine bekannte Wahrheit geprüft werden können.

BAU-43 verwendet dieselben Begriffe (Baugenehmigung, Fassade, Brandschutz) mit
gegenteiligem Stand. Taucht in Ergebnissen für BAU-42 "Baugenehmigung erteilt" auf,
wurden Projekte vermischt.

Aufruf: uv run python scripts/generate_sample_documents.py [--out data/raw]
"""

import argparse
import re
import zipfile
from datetime import UTC, datetime
from pathlib import Path

import docx
import openpyxl
import pymupdf
from docx.document import Document as DocxDocument
from openpyxl.styles import Font

PROJECT_ROOT = Path(__file__).resolve().parents[1]
FIXED_TIMESTAMP = datetime(2026, 9, 15, 12, 0, 0, tzinfo=UTC)

PDF_CSS = """
* { font-family: sans-serif; font-size: 10pt; }
h1 { font-size: 15pt; margin-bottom: 4pt; }
h2 { font-size: 12pt; margin-top: 10pt; }
p.meta { color: #555555; font-size: 9pt; }
table { border-collapse: collapse; width: 100%; }
th, td { border: 1px solid #999999; padding: 3pt; text-align: left; vertical-align: top; }
th { background-color: #e6e6e6; }
"""

Row = tuple[str, ...]


# --------------------------------------------------------------------------- Hilfsfunktionen


def html_table(header: Row, rows: list[Row]) -> str:
    head = "".join(f"<th>{cell}</th>" for cell in header)
    body = "".join(
        "<tr>" + "".join(f"<td>{cell}</td>" for cell in row) + "</tr>" for row in rows
    )
    return f"<table><tr>{head}</tr>{body}</table>"


def normalize_zip(path: Path) -> None:
    """DOCX/XLSX sind ZIP-Archive mit Zeitstempeln. Fixieren, damit gleicher Inhalt denselben SHA-256 ergibt."""
    with zipfile.ZipFile(path) as source:
        entries = [(info.filename, source.read(info)) for info in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, data in entries:
            if name == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:modified[^>]*>)[^<]*",
                    rb"\g<1>2026-09-15T12:00:00Z",
                    data,
                )
            info = zipfile.ZipInfo(name, date_time=FIXED_TIMESTAMP.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, data)


def write_pdf(path: Path, title: str, pages: list[str]) -> None:
    """Jeder Eintrag in `pages` wird genau eine PDF-Seite (stabile Seitenzahlen für Quellenbelege)."""
    doc = pymupdf.open()
    for number, html in enumerate(pages, start=1):
        page = doc.new_page(width=595, height=842)  # A4
        spare_height, _ = page.insert_htmlbox(
            page.rect + (56, 56, -56, -56), html, css=PDF_CSS, scale_low=1
        )
        if spare_height < 0:
            raise ValueError(
                f"{path.name}: Inhalt von Seite {number} passt nicht auf eine Seite"
            )
    pdf_date = FIXED_TIMESTAMP.strftime("D:%Y%m%d%H%M%S")
    doc.set_metadata(
        {
            "title": title,
            "author": "Synthetische Testdaten",
            "creationDate": pdf_date,
            "modDate": pdf_date,
        }
    )
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path, garbage=4, deflate=True, no_new_id=True)
    doc.close()


def write_protocol(
    path: Path,
    project: str,
    number: int,
    date: str,
    participants: list[str],
    tops: list[Row],
    next_meeting: str,
) -> None:
    """Baubesprechungsprotokoll als DOCX mit TOP-Tabelle (TOP | Thema | Sachstand / Beschluss | Verantwortlich | Termin)."""
    document: DocxDocument = docx.Document()
    document.add_heading(f"Protokoll Baubesprechung Nr. {number}", level=1)
    document.add_paragraph(f"Projekt: {project}")
    document.add_paragraph(f"Datum: {date}    Ort: Baucontainer, Besprechungsraum")
    document.add_heading("Teilnehmer", level=2)
    for participant in participants:
        document.add_paragraph(participant, style="List Bullet")
    document.add_heading("Tagesordnungspunkte", level=2)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for cell, text in zip(
        table.rows[0].cells,
        ("TOP", "Thema", "Sachstand / Beschluss", "Verantwortlich", "Termin"),
    ):
        cell.text = text
    for top in tops:
        for cell, text in zip(table.add_row().cells, top):
            cell.text = text
    document.add_paragraph()
    document.add_paragraph(f"Nächste Baubesprechung: {next_meeting}")
    document.add_paragraph(
        "Einwände gegen dieses Protokoll sind innerhalb von 5 Werktagen schriftlich mitzuteilen."
    )

    properties = document.core_properties
    properties.author = "Projektsteuerung"
    properties.title = f"Baubesprechung Nr. {number}"
    properties.created = properties.modified = FIXED_TIMESTAMP
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    normalize_zip(path)


# --------------------------------------------------------------------------- BAU-42

BAU42 = "BAU-42 – Wohnquartier Lindenhof, Haus A und B"
BAU42_PARTICIPANTS = [
    "Fr. Albers – Bauherr, Lindenhof Wohnbau GmbH",
    "Hr. Vogt – Projektsteuerung",
    "Hr. Yilmaz – Architekt, Nordlicht Architekten",
    "Hr. Kessler – Generalunternehmer, Kessler Bau GmbH",
    "Fr. Brandt – Fassadenplanung, Brandt Fassadentechnik",
]


def bau42_protocols(root: Path) -> None:
    write_protocol(
        root / "protokolle" / "2026-09-01_baubesprechung.docx",
        BAU42,
        23,
        "01.09.2026",
        BAU42_PARTICIPANTS,
        [
            (
                "1",
                "Baufortschritt Rohbau Haus A",
                (
                    "Decke über 2. OG betoniert. Verzug ca. 1 Woche wegen verspäteter Lieferung Bewehrungsstahl. "
                    "GU prüft Aufholmaßnahmen."
                ),
                "Kessler Bau",
                "08.09.2026",
            ),
            (
                "2",
                "Baugenehmigung Haus B",
                (
                    "Antrag (Az. BA-2026-0587) am 15.06.2026 eingereicht. Bauaufsicht fordert mit Schreiben vom 20.08.2026 "
                    "einen Prüfbericht Brandschutz. Ein Prüfsachverständiger für Brandschutz ist noch nicht beauftragt. "
                    "Zuständigkeit für die Beauftragung (Bauherr oder Architekt) ist ungeklärt."
                ),
                "offen",
                "08.09.2026",
            ),
            (
                "3",
                "Fassadenplanung Haus A",
                "Werk- und Montageplanung Fassade wird von Brandt Fassadentechnik bis 05.09.2026 zugesagt.",
                "Brandt Fassadentechnik",
                "05.09.2026",
            ),
            (
                "4",
                "Baustelleneinrichtung",
                "Zweiter Kran für Haus B: Standort abgestimmt, Sondernutzungserlaubnis Straßenraum liegt vor. Erledigt.",
                "Kessler Bau",
                "erledigt",
            ),
            (
                "5",
                "Nachträge",
                "Nachtrag N-05 Baugrubenverbau (12.300 EUR netto) geprüft und beauftragt. Erledigt.",
                "Lindenhof Wohnbau",
                "erledigt",
            ),
        ],
        "08.09.2026, 10:00 Uhr",
    )
    write_protocol(
        root / "protokolle" / "2026-09-08_baubesprechung.docx",
        BAU42,
        24,
        "08.09.2026",
        BAU42_PARTICIPANTS,
        [
            (
                "1",
                "Baufortschritt Rohbau Haus A",
                (
                    "Verzug hat sich auf 10 Arbeitstage (2 Wochen) erhöht. Vorgeschlagene Samstagsarbeit als "
                    "Aufholmaßnahme wird vom Bauherrn wegen Mehrkosten abgelehnt."
                ),
                "Kessler Bau",
                "15.09.2026",
            ),
            (
                "2",
                "Brandschutzprüfung Haus B",
                (
                    "Weiterhin kein Verantwortlicher für die Beauftragung des Prüfsachverständigen Brandschutz. "
                    "Bauherr sieht den Architekten in der Pflicht, Architekt verweist auf seinen Vertrag, der diese "
                    "Leistung nicht enthält. Eskalation durch Projektsteuerung."
                ),
                "ungeklärt",
                "15.09.2026",
            ),
            (
                "3",
                "Fassadenplanung Haus A",
                (
                    "Zugesagter Termin 05.09.2026 nicht eingehalten. Grund: Umstellung auf nichtbrennbare Dämmung "
                    "aufgrund Brandschutzanforderung. Neuer Termin 19.09.2026. Der Fassadenbauer Metallbau Hofer "
                    "benötigt freigegebene Pläne spätestens am 26.09.2026 für die Fertigung."
                ),
                "Brandt Fassadentechnik",
                "19.09.2026",
            ),
            (
                "4",
                "Baugenehmigung Haus B",
                (
                    "Bauaufsicht bestätigt telefonisch: Die Bearbeitung des Antrags ruht, bis der Prüfbericht "
                    "Brandschutz vorliegt. Geplanter Baubeginn Haus B: 12.10.2026."
                ),
                "Lindenhof Wohnbau",
                "15.09.2026",
            ),
            (
                "5",
                "Arbeitsschutz",
                "Begehung durch SiGeKo ohne Beanstandungen. Erledigt.",
                "SiGeKo",
                "erledigt",
            ),
        ],
        "15.09.2026, 10:00 Uhr",
    )
    write_protocol(
        root / "protokolle" / "2026-09-15_baubesprechung.docx",
        BAU42,
        25,
        "15.09.2026",
        BAU42_PARTICIPANTS,
        [
            (
                "1",
                "Baufortschritt Rohbau Haus A",
                (
                    "Verzug unverändert 2 Wochen. Fertigstellung Rohbau Haus A verschiebt sich vom 30.10.2026 "
                    "auf den 13.11.2026."
                ),
                "Kessler Bau",
                "22.09.2026",
            ),
            (
                "2",
                "Fassade Haus A – Nachtrag",
                (
                    'Nachtrag N-07 "Umstellung Fassadendämmung auf nichtbrennbare Mineralwolle" über 48.500 EUR netto '
                    "am 12.09.2026 eingereicht. Prüfung durch den Bauherrn offen, keine Freigabe. Fassadenplanung "
                    "weiterhin in Bearbeitung, Termin 19.09.2026 laut Fr. Brandt unsicher."
                ),
                "Lindenhof Wohnbau / Brandt Fassadentechnik",
                "19.09.2026",
            ),
            (
                "3",
                "Brandschutzprüfung Haus B",
                (
                    "Noch immer kein Verantwortlicher benannt. Bauaufsicht hat mit Schreiben vom 11.09.2026 eine Frist "
                    "bis 30.09.2026 gesetzt. Projektsteuerung: Ohne Beauftragung bis 18.09.2026 ist die Frist nicht "
                    "zu halten."
                ),
                "ungeklärt",
                "18.09.2026",
            ),
            (
                "4",
                "Baubeginn Haus B",
                (
                    "Projektsteuerung bewertet den Baubeginn am 12.10.2026 als gefährdet, da die Baugenehmigung Haus B "
                    "nicht erteilt ist. Eine Verschiebung um 4 bis 6 Wochen wird für möglich gehalten "
                    "(Einschätzung Projektsteuerung, nicht bestätigt)."
                ),
                "Projektsteuerung",
                "22.09.2026",
            ),
            (
                "5",
                "Entwässerung",
                "Genehmigung Kanalanschluss liegt vor. Erledigt.",
                "Kessler Bau",
                "erledigt",
            ),
        ],
        "22.09.2026, 10:00 Uhr",
    )


def bau42_status_reports(root: Path) -> None:
    header = (
        "<h1>Projektstatusbericht {month}</h1><p class='meta'>Projekt "
        + BAU42
        + " · Stand {stand} · Projektsteuerung</p>"
    )

    write_pdf(
        root / "statusberichte" / "statusbericht_2026-08.pdf",
        "Projektstatusbericht August 2026",
        [
            header.format(month="August 2026", stand="31.08.2026")
            + "<h2>1. Zusammenfassung</h2>"
            "<p>Gesamtstatus: <b>GELB</b>.</p>"
            "<p>Der Rohbau Haus A schreitet voran, der Verzug beträgt 5 Arbeitstage und gilt als aufholbar. "
            "Der Bauantrag für Haus B ist eingereicht und befindet sich in Prüfung durch die Bauaufsicht. "
            "Die Kosten liegen im Budget.</p>"
            "<h2>2. Wesentliche Ereignisse im Berichtszeitraum</h2>"
            "<p>Baugrube und Gründung Haus A wurden abgeschlossen. Die Bauaufsicht hat mit Schreiben vom "
            "20.08.2026 Unterlagen zum Brandschutz für Haus B nachgefordert.</p>",
            header.format(month="August 2026", stand="31.08.2026")
            + "<h2>3. Termine und Meilensteine</h2>"
            + html_table(
                ("Meilenstein", "Plan", "Prognose", "Status"),
                [
                    (
                        "Fertigstellung Rohbau Haus A",
                        "30.10.2026",
                        "30.10.2026",
                        "im Plan (Verzug aufholbar)",
                    ),
                    ("Fassadenplanung Haus A", "05.09.2026", "05.09.2026", "im Plan"),
                    ("Baugenehmigung Haus B", "15.09.2026", "15.09.2026", "in Prüfung"),
                    ("Baubeginn Haus B", "12.10.2026", "12.10.2026", "im Plan"),
                ],
            ),
            header.format(month="August 2026", stand="31.08.2026")
            + "<h2>4. Kosten und Nachträge</h2>"
            + html_table(
                ("Position", "Betrag (EUR netto)"),
                [
                    ("Gesamtbudget", "18.400.000"),
                    ("Beauftragt", "11.200.000"),
                    ("Beauftragte Nachträge", "210.000"),
                    ("Nachtrag N-05 Baugrubenverbau (in Prüfung)", "12.300"),
                    ("Kostenprognose", "18.150.000"),
                ],
            )
            + "<p>Die Kostenprognose liegt innerhalb des Budgets. Budgetreserve: 250.000 EUR.</p>",
        ],
    )

    write_pdf(
        root / "statusberichte" / "statusbericht_2026-09.pdf",
        "Projektstatusbericht September 2026",
        [
            header.format(month="September 2026", stand="15.09.2026")
            + "<h2>1. Zusammenfassung</h2>"
            "<p>Gesamtstatus: <b>ROT</b> (Vormonat: GELB).</p>"
            "<p>Der Rohbau Haus A liegt zwei Wochen hinter dem Plan. Die Fassadenplanung Haus A ist verspätet. "
            "Die Baugenehmigung für Haus B ist nicht erteilt, weil der geforderte Prüfbericht Brandschutz fehlt. "
            "Für die Beauftragung des Prüfsachverständigen ist bislang niemand verantwortlich benannt. "
            "Der Baubeginn Haus B am 12.10.2026 ist dadurch gefährdet.</p>",
            header.format(month="September 2026", stand="15.09.2026")
            + "<h2>2. Termine und Genehmigungen</h2>"
            + html_table(
                ("Meilenstein", "Plan", "Prognose", "Status"),
                [
                    (
                        "Fertigstellung Rohbau Haus A",
                        "30.10.2026",
                        "13.11.2026",
                        "verzögert (+10 AT)",
                    ),
                    (
                        "Fassadenplanung Haus A",
                        "05.09.2026",
                        "19.09.2026 (unsicher)",
                        "verzögert",
                    ),
                    (
                        "Prüfbericht Brandschutz Haus B",
                        "15.09.2026",
                        "offen",
                        "überfällig, nicht beauftragt",
                    ),
                    ("Baugenehmigung Haus B", "15.09.2026", "offen", "Antrag ruht"),
                    ("Baubeginn Haus B", "12.10.2026", "offen", "gefährdet"),
                ],
            )
            + "<p>Die Bauaufsicht hat für die fehlenden Unterlagen eine Frist bis 30.09.2026 gesetzt.</p>",
            header.format(month="September 2026", stand="15.09.2026")
            + "<h2>3. Kosten und Nachträge</h2>"
            + html_table(
                ("Nachtrag", "Gegenstand", "Betrag (EUR netto)", "Status"),
                [
                    ("N-05", "Baugrubenverbau", "12.300", "beauftragt"),
                    (
                        "N-06",
                        "Zusätzliche Bodenplatte Aufzugsunterfahrt",
                        "7.900",
                        "beauftragt",
                    ),
                    (
                        "N-07",
                        "Umstellung Fassadendämmung auf nichtbrennbare Mineralwolle",
                        "48.500",
                        "offen, ungeprüft",
                    ),
                ],
            )
            + "<p>Kostenprognose: 18.222.000 EUR. Bei Beauftragung von N-07 reduziert sich die Budgetreserve "
            "von 178.000 EUR auf 129.500 EUR.</p>",
            header.format(month="September 2026", stand="15.09.2026")
            + "<h2>4. Offene Punkte</h2>"
            "<ul>"
            "<li>Benennung eines Verantwortlichen für die Beauftragung des Prüfsachverständigen Brandschutz.</li>"
            "<li>Prüfung und Entscheidung Nachtrag N-07 Fassade.</li>"
            "<li>Übergabe der freigegebenen Fassadenpläne an Metallbau Hofer bis 26.09.2026.</li>"
            "<li>Entscheidung über Aufholmaßnahmen Rohbau Haus A (Samstagsarbeit wurde abgelehnt).</li>"
            "</ul>",
        ],
    )


def bau42_permits(root: Path) -> None:
    authority = "<p class='meta'>Stadt Musterstadt · Bauaufsichtsamt · Rathausplatz 1 · 12345 Musterstadt</p>"

    write_pdf(
        root / "genehmigungen" / "baugenehmigung.pdf",
        "Teilbaugenehmigung Haus A",
        [
            authority + "<h1>Teilbaugenehmigung</h1>"
            "<p class='meta'>Aktenzeichen BA-2026-0412 · Datum 28.04.2026</p>"
            "<p>Bauherr: Lindenhof Wohnbau GmbH</p>"
            "<p>Bauvorhaben: Neubau Wohnquartier Lindenhof, <b>Haus A</b> (Gebäudeklasse 4), Lindenstraße 12</p>"
            "<h2>Entscheidung</h2>"
            "<p>Für das oben genannte Bauvorhaben wird die Teilbaugenehmigung für Baugrube, Gründung und "
            "Rohbau von Haus A erteilt.</p>"
            "<p>Diese Teilbaugenehmigung umfasst <b>nicht</b> Haus B. Über den Bauantrag für Haus B "
            "(Az. BA-2026-0587) wird gesondert entschieden.</p>",
            authority
            + "<h2>Auflagen und Hinweise</h2>"
            + html_table(
                ("Nr.", "Auflage"),
                [
                    (
                        "A1",
                        "Der Baubeginn ist mindestens eine Woche vorher schriftlich anzuzeigen.",
                    ),
                    (
                        "A2",
                        (
                            "Vor Beginn von Ausbau und Fassadenarbeiten ist der Brandschutznachweis durch einen "
                            "Prüfsachverständigen für Brandschutz zu prüfen und der Prüfbericht vorzulegen."
                        ),
                    ),
                    (
                        "A3",
                        "Für die Außenwanddämmung sind ausschließlich nichtbrennbare Baustoffe zu verwenden.",
                    ),
                    (
                        "A4",
                        "Der Standsicherheitsnachweis wurde geprüft. Der Prüfbericht liegt vor.",
                    ),
                ],
            )
            + "<p>Gegen diesen Bescheid kann innerhalb eines Monats nach Bekanntgabe Widerspruch erhoben werden.</p>",
        ],
    )

    write_pdf(
        root / "genehmigungen" / "behördenkorrespondenz.pdf",
        "Behördenkorrespondenz Bauantrag Haus B",
        [
            authority
            + "<p class='meta'>An: Lindenhof Wohnbau GmbH · Datum 20.08.2026 · Az. BA-2026-0587</p>"
            "<h1>Bauantrag Haus B – Unvollständige Unterlagen</h1>"
            "<p>Sehr geehrte Damen und Herren,</p>"
            "<p>bei der Prüfung Ihres Bauantrags für Haus B haben wir festgestellt, dass folgende Unterlagen fehlen:</p>"
            "<ul><li>Prüfbericht eines Prüfsachverständigen für Brandschutz zum Brandschutznachweis,</li>"
            "<li>Fassadenschnitte mit Angabe der Dämmstoffe und deren Baustoffklasse.</li></ul>"
            "<p>Bis zur Vorlage dieser Unterlagen kann über den Antrag nicht entschieden werden.</p>"
            "<p>Mit freundlichen Grüßen<br>Bauaufsichtsamt</p>",
            (
                "<p class='meta'>Lindenhof Wohnbau GmbH · An: Bauaufsichtsamt Musterstadt · Datum 27.08.2026 · "
                "Az. BA-2026-0587</p>"
                "<h1>Ihr Schreiben vom 20.08.2026</h1>"
                "<p>Sehr geehrte Damen und Herren,</p>"
                "<p>die Beauftragung eines Prüfsachverständigen für Brandschutz ist in Vorbereitung. Die fehlenden "
                "Unterlagen werden wir voraussichtlich bis zum 15.09.2026 nachreichen.</p>"
                "<p>Mit freundlichen Grüßen<br>Fr. Albers, Lindenhof Wohnbau GmbH</p>"
            ),
            authority
            + "<p class='meta'>An: Lindenhof Wohnbau GmbH · Datum 11.09.2026 · Az. BA-2026-0587</p>"
            "<h1>Erinnerung und Fristsetzung</h1>"
            "<p>Sehr geehrte Damen und Herren,</p>"
            "<p>die mit Schreiben vom 20.08.2026 angeforderten Unterlagen liegen uns weiterhin nicht vor. "
            "Wir setzen Ihnen eine Frist bis zum <b>30.09.2026</b>. Nach fruchtlosem Fristablauf wird die "
            "Bearbeitung des Antrags eingestellt.</p>"
            "<p>Wir weisen darauf hin, dass die Bearbeitungszeit nach Vollständigkeit der Unterlagen "
            "erfahrungsgemäß etwa vier Wochen beträgt.</p>"
            "<p>Mit freundlichen Grüßen<br>Bauaufsichtsamt</p>",
        ],
    )


def bau42_schedule(root: Path) -> None:
    path = root / "terminplan" / "bauzeitenplan.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Bauzeitenplan"
    sheet.append(
        [f"Bauzeitenplan {BAU42}", "", "", "", "", "", "", "", "Stand 15.09.2026"]
    )
    sheet.append([])
    header = [
        "ID",
        "Vorgang",
        "Bereich",
        "Verantwortlich",
        "Start Plan",
        "Ende Plan",
        "Ende Prognose",
        "Verzug (AT)",
        "Status",
    ]
    sheet.append(header)
    for cell in sheet[3]:
        cell.font = Font(bold=True)
    rows: list[list[str | int | None]] = [
        [
            "V01",
            "Baugrube und Gründung Haus A",
            "Tiefbau",
            "Kessler Bau",
            "04.05.2026",
            "26.06.2026",
            "26.06.2026",
            0,
            "abgeschlossen",
        ],
        [
            "V02",
            "Rohbau Haus A",
            "Rohbau",
            "Kessler Bau",
            "29.06.2026",
            "30.10.2026",
            "13.11.2026",
            10,
            "verzögert",
        ],
        [
            "V03",
            "Werk- und Montageplanung Fassade Haus A",
            "Fassadenplanung",
            "Brandt Fassadentechnik",
            "15.07.2026",
            "05.09.2026",
            "19.09.2026",
            10,
            "verzögert",
        ],
        [
            "V04",
            "Fertigung Fassadenelemente Haus A",
            "Fassade",
            "Metallbau Hofer",
            "28.09.2026",
            "20.11.2026",
            "offen",
            None,
            "gefährdet",
        ],
        [
            "V05",
            "Fassadenmontage Haus A",
            "Fassade",
            "Metallbau Hofer",
            "16.11.2026",
            "29.01.2027",
            "12.02.2027",
            10,
            "gefährdet",
        ],
        [
            "V06",
            "Prüfbericht Brandschutz Haus B",
            "Genehmigung",
            "nicht benannt",
            "24.08.2026",
            "15.09.2026",
            "offen",
            None,
            "überfällig",
        ],
        [
            "V07",
            "Baugenehmigung Haus B",
            "Genehmigung",
            "Lindenhof Wohnbau",
            "15.06.2026",
            "15.09.2026",
            "offen",
            None,
            "offen",
        ],
        [
            "V08",
            "Meilenstein Baubeginn Haus B",
            "Meilenstein",
            "Kessler Bau",
            "12.10.2026",
            "12.10.2026",
            "offen",
            None,
            "gefährdet",
        ],
        [
            "V09",
            "Rohbau Haus B",
            "Rohbau",
            "Kessler Bau",
            "12.10.2026",
            "26.03.2027",
            "offen",
            None,
            "gefährdet",
        ],
    ]
    for row in rows:
        sheet.append(row)
    for column, width in zip("ABCDEFGHI", (6, 42, 16, 24, 12, 12, 14, 12, 14)):
        sheet.column_dimensions[column].width = width

    workbook.properties.creator = "Projektsteuerung"
    workbook.properties.created = FIXED_TIMESTAMP
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    normalize_zip(path)


# --------------------------------------------------------------------------- BAU-43 (Kontrollprojekt)

BAU43 = "BAU-43 – Erweiterung Grundschule am Park"


def bau43_documents(root: Path) -> None:
    write_protocol(
        root / "protokolle" / "2026-09-10_baubesprechung.docx",
        BAU43,
        11,
        "10.09.2026",
        [
            "Hr. Peters – Bauherr, Stadt Musterstadt, Hochbauamt",
            "Fr. Lange – Architektin, Lange + Partner",
            "Hr. Hansen – Generalunternehmer, Hansen Hochbau GmbH",
        ],
        [
            (
                "1",
                "Baufortschritt Rohbau",
                "Rohbau Erweiterungsbau im Plan. Richtfest am 02.10.2026 bestätigt.",
                "Hansen Hochbau",
                "02.10.2026",
            ),
            (
                "2",
                "Fassade",
                "Fassadenplanung am 01.09.2026 freigegeben. Fertigung läuft planmäßig.",
                "Lange + Partner",
                "erledigt",
            ),
            (
                "3",
                "Brandschutzprüfung",
                "Prüfsachverständiger Dr. Keller ist beauftragt. Prüfbericht Brandschutz liegt seit 20.07.2026 vor.",
                "Hochbauamt",
                "erledigt",
            ),
            (
                "4",
                "Aufzug",
                (
                    "Hersteller meldet Lieferverzug der Aufzugsanlage um 6 Wochen. Montage verschiebt sich auf Januar 2027. "
                    "Nachtrag N-02 für Zwischenlagerung (3.800 EUR netto) eingereicht, Prüfung offen."
                ),
                "Hansen Hochbau",
                "24.09.2026",
            ),
        ],
        "24.09.2026, 09:00 Uhr",
    )

    header = (
        "<h1>Projektstatusbericht September 2026</h1><p class='meta'>Projekt "
        + BAU43
        + " · Stand 15.09.2026</p>"
    )
    write_pdf(
        root / "statusberichte" / "statusbericht_2026-09.pdf",
        "Projektstatusbericht September 2026",
        [
            header + "<h2>1. Zusammenfassung</h2>"
            "<p>Gesamtstatus: <b>GELB</b>.</p>"
            "<p>Die Baugenehmigung ist seit 02.03.2026 vollständig erteilt. Rohbau und Fassade liegen im Plan. "
            "Einziger wesentlicher Punkt ist der Lieferverzug der Aufzugsanlage um 6 Wochen.</p>",
            header
            + "<h2>2. Termine und Kosten</h2>"
            + html_table(
                ("Meilenstein", "Plan", "Prognose", "Status"),
                [
                    ("Richtfest", "02.10.2026", "02.10.2026", "im Plan"),
                    ("Fassade fertig", "27.11.2026", "27.11.2026", "im Plan"),
                    ("Aufzugsmontage", "30.11.2026", "11.01.2027", "verzögert"),
                    ("Übergabe an Schulbetrieb", "15.03.2027", "15.03.2027", "im Plan"),
                ],
            )
            + "<p>Kosten: Budget 6.900.000 EUR, Prognose 6.780.000 EUR. Offener Nachtrag N-02 (3.800 EUR).</p>",
        ],
    )

    write_pdf(
        root / "genehmigungen" / "baugenehmigung.pdf",
        "Baugenehmigung Erweiterung Grundschule am Park",
        [
            (
                "<p class='meta'>Stadt Musterstadt · Bauaufsichtsamt · Rathausplatz 1 · 12345 Musterstadt</p>"
                "<h1>Baugenehmigung</h1>"
                "<p class='meta'>Aktenzeichen BA-2026-0133 · Datum 02.03.2026</p>"
                "<p>Bauherr: Stadt Musterstadt, Hochbauamt</p>"
                "<p>Bauvorhaben: Erweiterung Grundschule am Park, Parkweg 3</p>"
                "<h2>Entscheidung</h2>"
                "<p>Die Baugenehmigung wird vollständig erteilt. Der Prüfbericht Brandschutz und der geprüfte "
                "Standsicherheitsnachweis liegen vor. Die Fassade mit nichtbrennbarer Dämmung entspricht den "
                "Anforderungen.</p>"
            ),
        ],
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument(
        "--out",
        type=Path,
        default=PROJECT_ROOT / "data" / "raw",
        help="Zielverzeichnis",
    )
    args = parser.parse_args()

    bau42 = args.out / "BAU-42"
    bau42_protocols(bau42)
    bau42_status_reports(bau42)
    bau42_permits(bau42)
    bau42_schedule(bau42)
    bau43_documents(args.out / "BAU-43")

    for path in sorted(args.out.rglob("*.*")):
        if path.name != ".gitkeep":
            print(path.relative_to(args.out))


if __name__ == "__main__":
    main()
